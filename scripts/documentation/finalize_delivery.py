from pathlib import Path
import json,hashlib,html,urllib.parse,shutil
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'docs/professional-delivery';NAVY='18354B';TEAL='087F8C'
evidence=OUT/'00-阅读入口与复核/证据快照'
for name in ['audit/validation-latest.json','audit/display-fixes-v0.6.4.json','audit/reference-coverage.json','audit/official-runtime.json','audit/dice-layout-fixtures.json','audit/deck-copy-bounds-v0.6.4.json','assets/reference/SOURCES.json']:
 p=evidence/name;p.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(ROOT/name,p)
for src in (ROOT/'audit/test-logs/v0.6.4').glob('*.log'):
 dst=evidence/'audit/test-logs/v0.6.4'/src.name;dst.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(src,dst)
# Final delivery inventory is a native workbook, not a rendered Markdown table.
wb=Workbook();intro=wb.active;intro.title='使用说明';intro.append(['CASHFLOW 专业文档交付物清单']);intro.append(['基线','v0.6.4 / 2026-10-27']);intro.append(['用途','先按阶段选文件；DOCX可编辑，PDF固定阅读，XLSX用于协作台账，JSON用于契约。']);intro.append(['状态','本次补建资料并复核文件，未代替规则评审或用户签收。']);intro.append(['图例','R执行 / A负责 / C协商 / I知会；待执行与有历史证据分别记录。'])
ws=wb.create_sheet('交付物清单');ws.append(['原生文件与复核资料']);ws.append(['可按阶段、格式与用途筛选；历史日志不计为今日重跑。']);ws.append(['证据快照保留原始路径；批准和用户签收仍待实际执行。']);ws.append(['交付ID','阶段','文件名称','格式','用途','负责角色','状态','相对路径'])
role={'00':'技术／QA','01':'产品／制作人','02':'产品／策划／数值','03':'UX／UI／前端','04':'架构／前后端','05':'制作人／研发','06':'QA','07':'发布／运维','08':'产品／规则审阅／用户','09':'产品／研发／QA'}
purpose={'.docx':'原生可编辑规格／报告','.pdf':'固定排版评审副本','.xlsx':'原生数据／追踪工作簿','.pptx':'跨部门评审简报','.svg':'可编辑矢量图','.png':'图示渲染副本','.html':'本地阅读入口／设计原型','.json':'来源／契约／机器校验结果','.log':'既有执行证据','.txt':'阅读与复核说明'}
files=[p for p in sorted(OUT.rglob('*')) if p.is_file() and '证据快照' not in p.parts and p.name not in ['交付物清单.xlsx']]
# Include the known final entry and inventory itself without a circular hash claim.
for extra in [OUT/'开始阅读.html',OUT/'00-阅读入口与复核/交付物清单.xlsx']:
 if extra not in files:files.append(extra)
for i,p in enumerate(files,1):
 rel=p.relative_to(OUT);stage=rel.parts[0] if len(rel.parts)>1 else '00-阅读入口与复核';suffix=p.suffix.lower();ws.append([f'DEL-{i:03d}',stage,p.name,suffix[1:].upper(),purpose.get(suffix,'交付资料'),role.get(stage[:2],'技术／QA'),'待评审；证据按范围使用',str(rel)])
style_count=0
def sheet(name,heads,rows):
 global style_count
 w=wb.create_sheet(name);w.append([name]);w.append(['来源与事实范围明确；不预填批准。']);w.append(['筛选字段以当前工作基线为准。']);w.append(heads)
 for r in rows:w.append(r)
 return w
sheet('格式与阶段',['阶段','交付物性质','原生格式','接收部门'],[['01','章程／责任／风险','DOCX/PDF/XLSX','产品／制作人'],['02','PRD/GDD／规则数据／决策','DOCX/PDF/XLSX','策划／研发／QA'],['03','页面／标注／流程与原型','DOCX/PDF/XLSX/HTML/SVG','UX/UI／前端'],['04','架构／API／存档契约','DOCX/PDF/JSON/XLSX/SVG','前后端／QA'],['05','任务／追踪／交接','DOCX/PDF/XLSX','制作人／研发'],['06','用例／环境／缺陷／报告','DOCX/PDF/XLSX/JSON','QA'],['07','发布／配置／回退','DOCX/PDF/XLSX','发布／运维'],['08','UAT／遗留／签收','DOCX/PDF/XLSX','产品／用户'],['09','问题／变更／决策','DOCX/PDF/XLSX','产品／研发／QA']])
sources=json.loads((OUT/'00-阅读入口与复核/公开研究来源.json').read_text());sheet('公开研究来源',['资料','URL','访问日期','读取状态','使用边界'],[[r['title'],r['url'],r.get('accessed_at','2026-10-27'),r['status'],r.get('use','官方公开仓库入口；README不等于已审阅全部技术规范')] for r in sources])
sheet('阶段批准与缺口',['阶段','资料状态','执行／批准状态','下一步'],[['01-02','已形成章程、PRD、GDD和规则主表','范围／规则冲突待正式评审','指定比较版本与责任人'],['03-04','已形成设计、图、接口与Schema','逐页批准及契约业务评审待完成','关闭未知坐标与异常规格缺口'],['05-06','有实现与历史测试；本次复核文档契约','实际file、兼容、压测、安全仍待执行','按风险矩阵补实测'],['07-08','有阶段包、检查表与UAT方案','回退演练和最终签收未完成','锁定包SHA后用户验收'],['09','有历史修复与新台账','SLA及实际人员未指定','持续变更和回归']])
sheet('非功能建议目标',['指标ID','建议目标（非承诺）','测量条件','通过判断','批准状态'],[['NFR-P01','桌面冷启动可交互P95 ≤3秒','4核8GB桌面；稳定版Chrome；DPR1；1280×960；本地入口；20次，不含下载／解压','计时点从进入标题页到房间列表可操作；需先确认入口与设备','建议值，待技术／产品批准'],['NFR-P02','在线动作响应P95 ≤500ms，系统失败率≤1%','同一局域网；10房×2会话；持续10分钟；记录服务CPU／内存／写入','HTTP响应延迟；合法业务拒绝与系统失败分开统计','建议值，待技术／QA批准；未压测'],['NFR-R01','网络恢复后同会话恢复≤5秒','同浏览器会话；恢复网络作为计时起点；保留存档','恢复后状态与权威服务一致，无重复结算','建议值，待评审／实测'],['NFR-C01','关键数值字段允许范围内零溢出／重叠','指定字体、默认与放大比例、允许最大金额及持续开关','可读且可操作；结合DOM边界及人工截图','局部自动化已有；完整矩阵待批准'],['NFR-D01','正常保存恢复不丢状态；写入失败不静默成功','本机存储、在线磁盘、刷新与故障注入分开测','确认后的状态与恢复结果一致；有失败提示','候选准则，待完整演练'],['NFR-A01','先批准可访问性与原版像素目标的取舍','确定键盘、焦点、读屏、色彩范围；不直接宣称WCAG符合','按批准范围测试并记录例外','范围待评审']])
dash=sheet('交付摘要',['指标','工作簿公式／数值','解释'],[['DOCX数量','=COUNTIF(交付物清单!D5:D1000,"DOCX")','原生评审规格'],['PDF数量','=COUNTIF(交付物清单!D5:D1000,"PDF")','固定排版副本'],['XLSX数量','=COUNTIF(交付物清单!D5:D1000,"XLSX")','协作台账与真实数据'],['PPTX数量','=COUNTIF(交付物清单!D5:D1000,"PPTX")','跨部门简报'],['已取得最终签收',0,'未签收，不据历史自动化填通过'],['已有单元测试',353,'v0.6.4历史记录'],['已有浏览器检查',27,'v0.6.4历史记录'],['本轮HTTP契约样例',20,'本次实际校验；非重跑全部游戏测试']])
for w in wb:
 w.sheet_view.showGridLines=False
 if w.title=='使用说明':
  w.column_dimensions['A'].width=26;w.column_dimensions['B'].width=100
  for row in w:
   for c in row:c.font=Font(name='Microsoft YaHei',size=11,color=NAVY);c.alignment=Alignment(wrap_text=True,vertical='top')
  for i in range(1,6):w.row_dimensions[i].height=44
  continue
 n=w.max_column
 for i in [1,2,3]:
  w.merge_cells(start_row=i,start_column=1,end_row=i,end_column=n);c=w.cell(i,1);c.font=Font(name='Microsoft YaHei',size=17 if i==1 else 10,bold=i==1,color='FFFFFF' if i==1 else NAVY);c.fill=PatternFill('solid',fgColor=NAVY if i==1 else 'EAF3F5');c.alignment=Alignment(wrap_text=True,vertical='center');w.row_dimensions[i].height=34
 for c in w[4]:c.fill=PatternFill('solid',fgColor=TEAL);c.font=Font(name='Microsoft YaHei',bold=True,color='FFFFFF',size=10)
 for row in w.iter_rows(min_row=5):
  for c in row:c.font=Font(name='Microsoft YaHei',size=10,color=NAVY);c.alignment=Alignment(wrap_text=True,vertical='top')
 for col in range(1,n+1):w.column_dimensions[w.cell(4,col).column_letter].width=38 if col!=1 else 24
 if w.title=='公开研究来源':w.column_dimensions['B'].width=85
 if w.title=='交付物清单':w.column_dimensions['H'].width=90
 for i in range(5,w.max_row+1):w.row_dimensions[i].height=62
 w.freeze_panes='B5';w.auto_filter.ref=f'A4:{w.cell(w.max_row,n).coordinate}';w.print_title_rows='1:4';w.sheet_properties.pageSetUpPr.fitToPage=True;w.page_setup.orientation='landscape';w.page_setup.fitToWidth=1;w.page_setup.fitToHeight=0
wb.save(OUT/'00-阅读入口与复核/交付物清单.xlsx')
# Rich local entry: no remote fonts, JavaScript libraries, tracking or API requests.
blocks=[]
for folder in sorted(p for p in OUT.iterdir() if p.is_dir()):
 entries=[]
 for p in sorted(folder.iterdir()):
  if not p.is_file():continue
  rel=p.relative_to(OUT);suffix=p.suffix.lower();label=purpose.get(suffix,'交付资料')
  entries.append(f'<a class="file" href="{urllib.parse.quote(str(rel))}"><span class="tag">{html.escape(suffix[1:].upper())}</span><strong>{html.escape(p.name)}</strong><small>{html.escape(label)}</small></a>')
 blocks.append(f'<section id="s{folder.name[:2]}"><h2>{html.escape(folder.name)}</h2><div class="grid">'+''.join(entries)+'</div></section>')
nav=''.join(f'<a href="#s{p.name[:2]}">{html.escape(p.name)}</a>' for p in sorted(OUT.iterdir()) if p.is_dir())
page='''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Cashflow 专业开发文档交付</title><style>*{box-sizing:border-box}body{margin:0;background:#f2f6fa;color:#18354b;font:15px/1.7 Arial,"Microsoft YaHei",sans-serif}header{padding:42px max(24px,6vw);background:#18354b;color:white}h1{margin:6px 0;font-size:34px}header p{max-width:1000px}main{max-width:1260px;margin:auto;padding:26px}nav{display:flex;flex-wrap:wrap;gap:8px;margin:20px 0}nav a{padding:7px 12px;background:white;color:#087f8c;border:1px solid #cadce2;border-radius:6px;text-decoration:none}section{margin:32px 0}h2{border-bottom:2px solid #087f8c;padding-bottom:8px}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(300px,1fr));gap:12px}.file{display:grid;grid-template-columns:62px 1fr;gap:4px 10px;padding:15px;background:white;border:1px solid #d8e4eb;border-radius:8px;color:#18354b;text-decoration:none;overflow-wrap:anywhere}.file:hover{border-color:#087f8c}.file small{grid-column:2;color:#64748b}.tag{font-weight:bold;color:#087f8c}.note{padding:18px;background:#fff3dc;border-left:5px solid #b76b16}.proof{padding:20px;background:white;border:1px solid #d8e4eb;border-radius:8px}a{color:#087f8c}</style><header><div>ENGINEERING REVIEW · v0.6.4 · 2026-10-27</div><h1>Cashflow 专业开发文档交付</h1><p>按阶段组织的原生Office、PDF、图示、交互原型与机器契约。先读阶段复核报告，再按部门打开对应文件。建议用Office或LibreOffice编辑，PDF用于固定阅读。</p></header><main><div class="note"><strong>待评审基线，不代填批准。</strong> 现有实现、历史自动化证据、本次文件／契约复核和待执行的人工验收分开记录。尚未完成真实file入口的自动化验证，也未取得原版全部内容和多骰逐页位置证据。</div><nav>'''+nav+'''</nav><div class="proof"><strong>快速入口</strong><p><a href="00-%E9%98%85%E8%AF%BB%E5%85%A5%E5%8F%A3%E4%B8%8E%E5%A4%8D%E6%A0%B8/%E4%BA%A4%E4%BB%98%E4%BE%9D%E6%8D%AE%E4%B8%8E%E9%98%B6%E6%AE%B5%E5%A4%8D%E6%A0%B8%E6%8A%A5%E5%91%8A.pdf">先阅读阶段复核报告 PDF</a> · <a href="00-%E9%98%85%E8%AF%BB%E5%85%A5%E5%8F%A3%E4%B8%8E%E5%A4%8D%E6%A0%B8/%E4%BA%A4%E4%BB%98%E7%89%A9%E6%B8%85%E5%8D%95.xlsx">打开交付物清单 XLSX</a></p><p>原始日志与资料副本位于“00-阅读入口与复核／证据快照”；当前源码的指纹见同目录JSON。最终文件与契约检查结果见“复核结果.json”。</p></div>'''+''.join(blocks)+'''<section><h2>复核结论与边界</h2><p>文档可用性检查与游戏验收不同：本次核对原生文件结构、公式、资料数量、接口样例和PDF阅读效果。353单元测试及27浏览器检查来自既有v0.6.4报告，未伪装成今日重跑。用户签收、性能和兼容目标仍需实际责任人确认和执行。</p><a href="00-%E9%98%85%E8%AF%BB%E5%85%A5%E5%8F%A3%E4%B8%8E%E5%A4%8D%E6%A0%B8/%E5%A4%8D%E6%A0%B8%E7%BB%93%E6%9E%9C.json">查看机器复核结果</a></section></main></html>'''
(OUT/'开始阅读.html').write_text(page)
(OUT/'00-阅读入口与复核/阅读与复现说明.txt').write_text('''Cashflow 专业文档交付（v0.6.4实现基线，2026-10-27文档评审稿）

先解压，打开最外层“开始阅读.html”。Word／Excel／PowerPoint可编辑，PDF固定排版，HTML是设计原型，JSON是接口及数据契约。不是游戏安装包。

批准状态保持未签署。公开来源中失败的访问不当作已读标准；企业公开示例不是内部文档模板。

本次复核工具位于仓库scripts/documentation：build_delivery.py生成原生资料，draw_architecture.py生成图，render_diagrams.cjs渲染图，style_delivery.py增强工作簿及简报，check_api_examples.cjs验证项目自身接口，finalize_delivery.py建立阅读入口和清单。依赖python-docx、openpyxl、python-pptx、jsonschema、PyMuPDF、Playwright与LibreOffice。

补充复现顺序：生成基础资料→draw_architecture→render_diagrams→style_delivery→add_data_dictionary→add_rule_review→polish_contract→Office转PDF并重算工作簿→接口和入口校验→finalize_delivery→apply_metadata→check_delivery。所有步骤以仓库脚本及校验结果为准，重新生成后需再次复核。

建议用PDF进行固定版式评审；Office编辑时若本机缺少Noto Sans CJK字体，可能发生字体回退。PDF已实际导出并检查中文。

PDF由LibreOffice实际导出；Excel公式由LibreOffice重算。脚本不代签验收，不替用户执行file入口和完整UAT。质量结果及哈希文件随包保存。
''')
print('Created native inventory workbook, evidence snapshots and local reading portal.')
