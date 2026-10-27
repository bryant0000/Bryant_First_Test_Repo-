from pathlib import Path
from docx import Document
from docx.shared import Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook
from openpyxl.styles import Font,Alignment
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'docs/professional-delivery'
rows=[['初始现金','参考profile：cash = savings','旧profile为savings + flow，不混用'],['月支出','tax + other + Σloan.payment + kids × childCost','school初始化兼容差异见DEC-04'],['内圈发薪','salary + Σasset.income − 月支出','本轮确认及经过格由规则处理'],['快车道进入','P = 内圈被动收入，要求P > 月支出','确认后fastBase = P × 100'],['入场现金／位置','cash替换为fastBase；pos=42','charity及skip清零；table自动交棒'],['快车道报表收入','fastBase + fastGain','报表展示值与发薪口径需分别审阅'],['当前快车道发薪','fastBase + fastGain + salary − ratExpenses','engine.totals当前参考profile行为；DEC-09待审阅'],['新增收入胜利','fastGain >= 50,000','以规则check触发，不依赖动画'],['梦想胜利','购买自己dreamId对应位置','费用结合dreamMultiplier'],['慈善','内圈总收入10%，至少1；快车道100,000','内圈后续3次可选1/2骰；快车道1/2/3骰'],['失业／破产恢复','失业先支付月支出，休整2次；重整恢复休整3次','休整确认后交棒，不能额外掷骰'],['银行借款','每次正整千，参考profile上限9,999,000','余额及现金增加，月供为银行余额10%'],['偿还银行债','整千舍入并受现金与余额约束','非银行个人债务全额结清'],['证券确认','买入1..99,999整数份；卖出≤持仓','买入扣price×q；卖出先处理低成本批次'],['投资检定','die >= rollmin为成功','先付投入；确认结果不重复扣款']]
p=OUT/'02-需求与游戏规则/GDD玩法与数值设计.docx';d=Document(p)
# Idempotent only on a newly built GDD; record marker avoids duplicate append on rerun.
if not any('关键数值口径表' in x.text for x in d.paragraphs):
 marker=d.add_paragraph('关键数值口径表（当前实现，待规则评审）',style='Heading 2');t=d.add_table(rows=1,cols=3);t.style='Table Grid'
 for c,v in zip(t.rows[0].cells,['机制','公式／行为','条件与评审边界']):c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=v
 for i,row in enumerate(t.rows):
  no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
  if i==0:
   flag=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(flag)
  for c in row.cells:
   if i==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'18354B');c._tc.get_or_add_tcPr().append(shade)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(5)
    for run in para.runs:run.font.name='Noto Sans CJK SC';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');run.font.size=Pt(9);run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '18354B');run.bold=i==0
 anchor=next(x for x in d.paragraphs if x.text=='事件、胜负与队列');anchor._p.addnext(t._tbl);anchor._p.addnext(marker._p);d.save(p)
new=['DEC-09','快车道发薪口径','当前参考profile发薪包含内圈salary − ratExpenses；报表收入仅fastBase + fastGain','对照原版实际发薪，确认目标口径后再决定是否改规则','规则审阅／产品','待核验']
for p in OUT.rglob('*.xlsx'):
 w=load_workbook(p);changed=False
 for title in ['规则决策','未决阻断']:
  if title in w:
   ws=w[title]
   if not any(ws.cell(i,1).value=='DEC-09' for i in range(5,ws.max_row+1)):
    ws.append(new);i=ws.max_row;ws.row_dimensions[i].height=80
    for c in ws[i]:c.font=Font(name='Microsoft YaHei',size=10);c.alignment=Alignment(wrap_text=True,vertical='top')
    for t in ws.tables.values():t.ref=f'A4:F{i}'
    ws.auto_filter.ref=f'A4:F{i}';changed=True
 if changed:w.save(p)
source=ROOT/'scripts/documentation/build_delivery.py';s=source.read_text();tag="risk=[('RSK-01'"
if "decisions.append(('DEC-09'" not in s:s=s.replace(tag,"decisions.append("+repr(tuple(new))+")\n"+tag)
source.write_text(s)
print('Added exact current numerical rules and DEC-09 review item; game source unchanged.')
