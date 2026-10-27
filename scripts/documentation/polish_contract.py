from pathlib import Path
import json,copy
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2];p=ROOT/'docs/professional-delivery/04-技术与接口契约';oas=json.loads((p/'openapi.json').read_text());g=json.loads((p/'game.schema.json').read_text());store=json.loads((p/'room-store.schema.json').read_text());sc=oas['components']['schemas']
g['properties']['ruleset']={'type':'string','const':'classic-research-v3','description':'当前v0.6.4数据修订；升级数据时需要新契约／迁移审查。'};g['$defs']['Player']['properties']['ruleset']=copy.deepcopy(g['properties']['ruleset']);g['$defs']['Pending']['properties']['previous']={'$ref':'#/$defs/Pending'}
def adapt(x,prefix):
 if isinstance(x,dict):return {k:adapt(v,prefix) for k,v in x.items() if k!='$id'}
 if isinstance(x,list):return [adapt(v,prefix) for v in x]
 if isinstance(x,str):return x.replace('#/$defs/',prefix)
 return x
sc['Game']=adapt(g,'#/components/schemas/Game/$defs/')
sc['RoomSummary']={'type':'object','required':['id','name','capacity','players','private','version'],'properties':{'id':{'type':'string'},'name':{'type':'string'},'capacity':{'type':'integer','minimum':1,'maximum':6},'players':{'type':'integer','minimum':0,'maximum':6},'private':{'type':'boolean'},'version':{'type':'integer'}}}
sc['SeatSummary']={'type':'object','required':['seat','name','ready','left','away'],'properties':{'seat':{'type':'integer','minimum':0,'maximum':5},'name':{'type':'string'},'ready':{'type':'boolean'},'left':{'type':'boolean'},'away':{'type':'boolean'}}}
sc['ChatMessage']={'type':'object','required':['id','author','message','kind'],'properties':{'id':{'type':'integer'},'author':{'type':'string'},'message':{'type':'string'},'kind':{'type':'string','description':'目前为chat或action；系统动作记录也使用action。'}}}
sc['RoomList']['properties']['rooms']['items']={'$ref':'#/components/schemas/RoomSummary'}
sc['RoomView']['properties']['players']['items']={'$ref':'#/components/schemas/SeatSummary'};sc['RoomView']['properties']['chat']['items']={'$ref':'#/components/schemas/ChatMessage'}
sc['SessionRoomView']={'allOf':[{'$ref':'#/components/schemas/RoomView'},{'type':'object','required':['token'],'properties':{'token':{'type':'string','minLength':1,'description':'创建／加入签发的房间会话；不输出到公开日志。'}}}]}
for path,code in [('/api/rooms','201'),('/api/rooms/{id}/join','200')]:oas['paths'][path]['post']['responses'][code]['content']['application/json']['schema']={'$ref':'#/components/schemas/SessionRoomView'}
store['$defs']={'Game':adapt(g,'#/$defs/Game/$defs/')};store['properties']['rooms']['items']['properties']['game']={'anyOf':[{'type':'null'},{'$ref':'#/$defs/Game'}]}
for name,data in [('game.schema.json',g),('room-store.schema.json',store),('openapi.json',oas)]:(p/name).write_text(json.dumps(data,ensure_ascii=False,indent=2))
f=p/'接口指令与数据字典.xlsx';w=load_workbook(f);ws=w['HTTP接口']
for row in range(5,ws.max_row+1):
 if ws.cell(row,1).value=='POST' and ws.cell(row,2).value in ['/api/rooms','/api/rooms/{id}/join']:ws.cell(row,4,'SessionRoomView（必须含token）')
w.save(f)
print('Strengthened room/seat/chat and issued-session response types; nested store games use the reviewed schema.')
# Reviewed client examples and a nonempty disk shape with deliberately unusable auth data.
game_example=json.loads((p/'game.example.json').read_text())
(p/'room-store.example.json').write_text(json.dumps({'version':1,'rooms':[{'id':'documentation-example','name':'文档样例（非生产存档）','capacity':1,'host':0,'version':1,'seats':[{'name':'契约样例','ready':True,'left':False,'away':False,'lastSeen':0,'tokenHash':'0'*64,'jobId':4}],'salt':'documentation-only','passwordHash':None,'chatRevision':0,'chat':[],'game':game_example}]},ensure_ascii=False,indent=2))
oas['paths']['/api/rooms']['post']['requestBody']['content']['application/json']['example']={'name':'演示房间','capacity':2,'playerName':'甲','password':'1234'}
oas['paths']['/api/rooms/{id}/action']['post']['requestBody']['content']['application/json']['example']={'version':5,'action':'roll','args':[1]}
(p/'openapi.json').write_text(json.dumps(oas,ensure_ascii=False,indent=2))
