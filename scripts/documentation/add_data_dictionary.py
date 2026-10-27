from pathlib import Path
import json,subprocess
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
ROOT=Path(__file__).resolve().parents[2];p=ROOT/'docs/professional-delivery/02-需求与游戏规则/需求规则与数值主表.xlsx'
D=json.loads(subprocess.check_output(['node','-e',"console.log(JSON.stringify(require('./reference-data.js')));"],cwd=ROOT))
labels={'id':'职业标识','name':'显示名','english':'原英文名','salary':'工资','tax':'税','other':'其他生活支出','savings':'一次性储蓄','childCost':'每孩月支出','sourceId':'来源记录ID','anyoneBuy':'共享报价可买入','key':'资产匹配键／特定模式增量','type':'牌／格类型','units':'单位／套数','symbol':'证券代码','price':'每份报价','dividend':'每份净收入','rangeMin':'常见报价下界','rangeMax':'常见报价上界','down':'首付','value':'资产总价','debt':'附带抵押贷款','income':'净现金流','factor':'拆并股倍率','mode':'市场效果类型','global':'是否共享市场效果','amount':'市场报价／增价／费用／倍率','minUnits':'匹配最少单位','cost':'卡面费用／投入','child':'子女适用条件','payment':'分期新增月供','position':'外圈位置','icon':'显示图标标记','artKey':'原始图集帧名','rollmin':'投资成功阈值','reward':'成功奖励现金'}
w=load_workbook(p)
for title in ['机制数据字典','数值不变量']:
 if title in w:del w[title]
ws=w.create_sheet('机制数据字典');ws.append(['机制数据字典｜当前字段逐项映射']);ws.append(['原始字段保留；金额与倍率按上下文区分，不能把所有amount都标为美元。']);ws.append(['本表由实际数据记录遍历生成，规则解释按engine.js分支复核。']);ws.append(['数据域','字段','中文含义','实际类型','单位／约束','例值'])
for domain in ['jobs','small','big','market','doodads','dreams','fast']:
 keys=list(dict.fromkeys(k for r in D[domain] for k in r))
 for k in keys:
  values=[r[k] for r in D[domain] if k in r];v=values[0];types='/'.join(sorted({type(x).__name__ for x in values}));unit='字符串／匹配与显示标识'
  label=labels.get(k)
  if k.endswith('Payment'):label=k[:-7]+'贷款月供';unit='金额／月'
  elif k.endswith('Debt'):label=k[:-4]+'贷款余额';unit='金额；参考profile初始化school债务另见DEC-04'
  elif k in ['salary','tax','other','childCost','dividend','payment']:unit='金额／月（内圈）'
  elif k in ['cost','down','value','debt','price','rangeMin','rangeMax','savings','reward']:unit='金额；price/range为每份，savings仅开局一次'
  elif k=='income':unit='金额／月（内圈资产）或金额／现金流日（快车道）'
  elif k=='amount':unit='market.mode: sell报价、selladd增价、fee费用、sellmultiply倍率；startup增量来自key；brother本次净收款按规则处理'
  elif k=='key' and domain=='market':unit='通常为资产匹配键；startup中为收入增量数值字符串'
  elif k in ['units','minUnits']:unit='计数；匹配时结合资产key和单位类型'
  elif k=='factor':unit='无量纲倍率；拆股2或并股0.5'
  elif k=='position':unit='0基外圈索引，0..47'
  elif k=='rollmin':unit='骰点成功阈值；0表示非风险'
  elif k=='child':unit='false无子女限制；has或per按引擎决定有孩或逐孩费用'
  elif k in ['anyoneBuy','global']:unit='布尔标记；不是在线座位权限'
  ws.append([domain,k,label or k,types,unit,json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list,bool)) else v])
ws2=w.create_sheet('数值不变量');ws2.append(['数值不变量与检查前提']);ws2.append(['仅在没有同时发生升阶／其他强制效果时比较单次现金变化。']);ws2.append(['实现应使用独立例值验证，不能只重写实现公式。']);ws2.append(['ID','情形','应满足关系','前提／测试依据'])
for row in [('INV-01','证券买入','现金减少price×q，持股增加q','合法整数份数；不足现金先走贷款，未确认不买入'),('INV-02','证券出售','现金增加price×q，按便宜批次扣除成本','持仓足够；分红按剩余数量重新计算'),('INV-03','银行借款','现金和银行余额加1000n，月供加100n','参考profile内圈；单位和借款上限经校验'),('INV-04','债务偿还','现金减少实际偿还额；余额／月供按类型减少','非银行全额；银行金额舍入和现金上限'),('INV-05','月净流入','工资＋资产收入−税−其他−月供−子女支出','参考profile开局省略school债务兼容差异'),('INV-06','房产资产栏','显示总价value，不把首付cost当总价','清算所得使用所选规则的首付基准'),('INV-07','取消与查看','未执行交易；草稿及卡片恢复','取消、确认、关闭窗口与结束回合分别测试')]:ws2.append(row)
for sheet in [ws,ws2]:
 for i in [1,2,3]:sheet.merge_cells(start_row=i,start_column=1,end_row=i,end_column=sheet.max_column);c=sheet.cell(i,1);c.font=Font(name='Microsoft YaHei',size=16 if i==1 else 10,color='FFFFFF' if i==1 else '18354B',bold=i==1);c.fill=PatternFill('solid',fgColor='18354B' if i==1 else 'EAF3F5');c.alignment=Alignment(wrap_text=True);sheet.row_dimensions[i].height=34
 for c in sheet[4]:c.fill=PatternFill('solid',fgColor='087F8C');c.font=Font(name='Microsoft YaHei',color='FFFFFF',bold=True)
 for col in range(1,sheet.max_column+1):sheet.column_dimensions[sheet.cell(4,col).column_letter].width=25 if col<5 else 60
 for row in sheet.iter_rows(min_row=5):
  for c in row:c.font=Font(name='Microsoft YaHei',size=10);c.alignment=Alignment(wrap_text=True,vertical='top')
  sheet.row_dimensions[row[0].row].height=54
 sheet.freeze_panes='B5';sheet.auto_filter.ref=f'A4:{sheet.cell(sheet.max_row,sheet.max_column).coordinate}';sheet.print_title_rows='1:4';sheet.page_setup.orientation='landscape';sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
w.save(p);print('Added actual data field dictionary and monetary invariants.')
