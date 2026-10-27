from pathlib import Path
import json,hashlib,re,zipfile,xml.etree.ElementTree as ET,warnings,subprocess
from collections import Counter
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook,Workbook
from openpyxl.styles import Font,PatternFill,Alignment
import fitz,jsonschema
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'docs/professional-delivery';rows=[];errors=[];warnings.filterwarnings('ignore',category=DeprecationWarning)
def record(path,kind,details,passed=True):
 rows.append({'file':str(path.relative_to(OUT)),'kind':kind,'passed':passed,'details':details})
 if not passed:errors.append(rows[-1])
result_path=OUT/'00-阅读入口与复核/复核结果.json'
if not result_path.exists():result_path.write_text(json.dumps({'status':'复核进行中，尚未发布'},ensure_ascii=False))
for p in sorted(OUT.rglob('*')):
 if not p.is_file() or '证据快照' in p.parts:continue
 try:
  if p.suffix in ['.docx','.xlsx','.pptx']:
   with zipfile.ZipFile(p) as z:
    assert z.testzip() is None
    for name in z.namelist():
     if name.endswith('.xml'):ET.fromstring(z.read(name))
   if p.suffix=='.docx':
    d=Document(p);assert len(d.tables)>=3,'Missing control, body design/check or review table';assert len(d.sections)>0;assert p.with_suffix('.pdf').exists();record(p,'DOCX',{'tables':len(d.tables),'pdf_copy':True})
   elif p.suffix=='.pptx':
    d=Presentation(p);assert len(d.slides)==15
    for s in d.slides:
     for sh in s.shapes:
      assert sh.left>=-1 and sh.top>=-1
      assert sh.left+sh.width<=d.slide_width+20000 and sh.top+sh.height<=d.slide_height+20000
    record(p,'PPTX',{'slides':len(d.slides),'shape_bounds':'pass'})
   else:
    w=load_workbook(p,data_only=False);v=load_workbook(p,data_only=True);formula_count=0
    for s in w:
     for row in s:
      for c in row:
       if c.data_type=='f':formula_count+=1;assert v[s.title][c.coordinate].value is not None,(s.title,c.coordinate)
       if v[s.title][c.coordinate].data_type=='e':raise AssertionError('Excel formula error '+s.title+' '+c.coordinate)
    record(p,'XLSX',{'sheets':len(w.sheetnames),'formulas_with_cache':formula_count})
  elif p.suffix=='.pdf':
   d=fitz.open(p);assert len(d)>0;bad=[];text=''
   for i,page in enumerate(d):
    text+=page.get_text();assert len(page.get_text().strip())>12
    for block in page.get_text('dict')['blocks']:
     for line in block.get('lines',[]):
      for s in line['spans']:
       x0,y0,x1,y1=s['bbox']
       if x0<-.5 or y0<-.5 or x1>page.rect.width+.5 or y1>page.rect.height+.5:bad.append({'page':i+1,'text':s['text'][:80],'bbox':s['bbox']})
   assert '\ufffd' not in text
   assert not bad,bad
   record(p,'PDF',{'pages':len(d),'text_characters':len(text),'out_of_page_spans':0})
  elif p.suffix in ['.json']:
   json.loads(p.read_text());record(p,'JSON',{'parse':'pass'})
  elif p.suffix=='.svg':ET.fromstring(p.read_text());record(p,'SVG',{'xml':'pass'})
 except Exception as e:record(p,p.suffix,{'error':str(e)},False)
# Data quantities and cached arithmetic versus the real rule engine.
data=OUT/'02-需求与游戏规则/需求规则与数值主表.xlsx';w=load_workbook(data,data_only=True)
for title,n in [('职业数据',12),('小买卖',55),('大买卖',40),('市场',45),('额外支出',42),('梦想',23),('快车道',48)]:assert w[title].max_row-4==n
for i in range(5,12):assert w['数量复核'].cell(i,5).value=='PASS'
facts=json.loads(subprocess.check_output(['node','-e',"const E=require('./engine.js');console.log(JSON.stringify(E.jobs.map((j,i)=>E.ratTotals(E.create(i,0,E.REFERENCE_PROFILE)))));"],cwd=ROOT))
for i,f in enumerate(facts,5):assert w['职业财务摘要'].cell(i,6).value==f['expenses'] and w['职业财务摘要'].cell(i,7).value==f['flow']
record(data,'DATA-CROSSCHECK',{'jobs':12,'mechanical_cards':182,'dreams':23,'fast_spaces':48,'role_formula_results':'match actual engine'})
uat=load_workbook(OUT/'08-用户验收/UAT签收与遗留事项.xlsx',data_only=True)
for i in range(5,13):assert '未签收' in uat['验收场景'].cell(i,7).value
for i in range(5,9):assert uat['签收记录'].cell(i,3).value=='未签署'
record(OUT/'08-用户验收/UAT签收与遗留事项.xlsx','APPROVAL-BOUNDARY',{'approval':'not fabricated','uat':'not signed'})
# Keep the docs tied to the unmodified implementation.
finger=json.loads((OUT/'00-阅读入口与复核/源码复核指纹.json').read_text())
for name,digest in finger['files'].items():assert hashlib.sha256((ROOT/name).read_bytes()).hexdigest()==digest
record(OUT/'00-阅读入口与复核/源码复核指纹.json','SOURCE-FINGERPRINT',{'files':len(finger['files']),'unchanged':True})
tech=OUT/'04-技术与接口契约';schema=json.loads((tech/'game.schema.json').read_text());store=json.loads((tech/'room-store.schema.json').read_text());jsonschema.Draft202012Validator.check_schema(schema);jsonschema.Draft202012Validator.check_schema(store)
jsonschema.validate(json.loads((tech/'game.example.json').read_text()),schema);jsonschema.validate(json.loads((tech/'room-store.example.json').read_text()),store)
example=json.loads((tech/'game.example.json').read_text());example['players'][0]['pos']=24
assert not jsonschema.Draft202012Validator(schema).is_valid(example)
oas=json.loads((tech/'openapi.json').read_text());assert oas['openapi']=='3.1.0';ids=[]
def pointers(v):
 if isinstance(v,dict):
  for k,x in v.items():
   if k=='$ref' and x.startswith('#/'):
    target=oas
    for part in x[2:].split('/'):target=target[part.replace('~1','/').replace('~0','~')]
   else:pointers(x)
 elif isinstance(v,list):
  for x in v:pointers(x)
pointers(oas)
for item in oas['paths'].values():
 for op in item.values():ids.append(op['operationId'])
assert len(ids)==len(set(ids))==10
record(tech/'openapi.json','CONTRACT-STRUCTURE',{'operations':10,'all_local_refs_resolved':True,'schema_examples':2,'negative_game_position':'rejected'})
api=json.loads((OUT/'00-阅读入口与复核/接口与Schema复核结果.json').read_text());assert api['validated_count']==20
# Relative links, images and local entry must exist. Fragment links are local headings.
for p in OUT.rglob('*.html'):
 for target in re.findall(r'(?:href|src)="([^"]+)"',p.read_text()):
  if target.startswith(('#','http','data:')):continue
  import urllib.parse
  q=p.parent/urllib.parse.unquote(target.split('#')[0]);assert q.exists(),(str(p),target)
record(OUT/'开始阅读.html','LOCAL-LINKS',{'relative_links':'all resolve'})
wb=Workbook();s=wb.active;s.title='文件复核结果';s.append(['Cashflow 文档文件层复核，不代替产品验收']);s.append(['文件','格式／检查','结果','复核信息'])
for row in rows:s.append([row['file'],row['kind'],'通过' if row['passed'] else '不通过',json.dumps(row['details'],ensure_ascii=False)])
s.freeze_panes='B3';s.auto_filter.ref=f'A2:D{s.max_row}'
for c in s[2]:c.fill=PatternFill('solid',fgColor='18354B');c.font=Font(name='Microsoft YaHei',bold=True,color='FFFFFF')
for col,width in [('A',75),('B',25),('C',12),('D',90)]:s.column_dimensions[col].width=width
for row in s.iter_rows(min_row=3):
 for c in row:c.font=Font(name='Microsoft YaHei',size=10);c.alignment=Alignment(wrap_text=True,vertical='top')
 s.row_dimensions[row[0].row].height=45
s2=wb.create_sheet('复核边界');s2.append(['事项','结论']);s2.append(['本轮HTTP契约样例',20]);s2.append(['历史游戏测试','353单测／27浏览器来自既有v0.6.4，不冒充重跑']);s2.append(['正式批准／用户签收','未签署']);s2.append(['尚未通过范围','真实file入口全矩阵、容量／性能、安全审计、原版全部内容一致性'])
s2.column_dimensions['A'].width=30;s2.column_dimensions['B'].width=110
qa=OUT/'00-阅读入口与复核/交付复核清单.xlsx';wb.save(qa);assert load_workbook(qa).sheetnames==['文件复核结果','复核边界']
result={'date':'2026-10-27','scope':'专业交付文件、实际代码数据、接口样例与资料一致性；不是游戏完整重新验收','passed':not errors,'counts':dict(Counter(p.suffix for p in OUT.rglob('*') if p.is_file() and '证据快照' not in p.parts)),'checks':rows,'errors':errors,'http_contract_exchanges':20,'manual_review':'抽查PRD、交付复核报告、UX正文、架构图和评审简报；完整PDF均做文字边界与解析检查','known_limits':['没有大公司内部模板或标准认证的声明','未重跑历史27套游戏浏览器检查','真实file入口、性能安全与最终UAT仍待执行','多骰原版运行坐标未全部核验']}
(OUT/'00-阅读入口与复核/复核结果.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps({'passed':not errors,'files_checked':len(rows),'errors':errors,'counts':result['counts']},ensure_ascii=False))
assert not errors
