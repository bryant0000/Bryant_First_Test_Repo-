from pathlib import Path
import json,copy
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart,Reference
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'docs/professional-delivery';NAVY='18354B';TEAL='087F8C'
# Add a human-oriented summary without changing the original wide source tables.
data=json.loads((OUT/'00-阅读入口与复核/构建事实快照.json').read_text())
facts=json.loads(__import__('subprocess').check_output(['node','-e',"const E=require('./engine.js');console.log(JSON.stringify(E.jobs.map((j,i)=>({name:j.name,...E.ratTotals(E.create(i,0,E.REFERENCE_PROFILE)),savings:j.savings,tax:j.tax,other:j.other}))));"],cwd=ROOT))
p=OUT/'02-需求与游戏规则/需求规则与数值主表.xlsx';wb=load_workbook(p)
if '职业财务摘要' in wb:del wb['职业财务摘要']
ws=wb.create_sheet('职业财务摘要',1);ws.append(['职业财务摘要｜参考profile开局']);ws.append(['公式只对应当前参考profile，教育贷款初始化兼容差异见DEC-04。']);ws.append(['金额为数值；工资与支出图用于设计复核，不代表难度平衡已获批准。']);ws.append(['职业','工资','税','其他支出','贷款月供','总支出（公式）','净流入（公式）','一次性储蓄'])
for i,j in enumerate(facts,5):ws.append([j['name'],j['salary'],j['tax'],j['other'],j['payments'],f'=C{i}+D{i}+E{i}',f'=B{i}-F{i}',j['savings']])
chart=BarChart();chart.title='职业工资与月支出';chart.y_axis.title='金额';chart.x_axis.title='职业';chart.add_data(Reference(ws,min_col=2,max_col=2,min_row=4,max_row=16),titles_from_data=True);chart.add_data(Reference(ws,min_col=6,max_col=6,min_row=4,max_row=16),titles_from_data=True);chart.set_categories(Reference(ws,min_col=1,min_row=5,max_row=16));chart.height=11;chart.width=23;ws.add_chart(chart,'J4');ws.freeze_panes='B5';ws.auto_filter.ref='A4:H16';wb.save(p)
# Native Office review styling and editable status choices.
for p in OUT.rglob('*.xlsx'):
 w=load_workbook(p)
 for ws in w:
  ws.sheet_view.showGridLines=False
  if ws.title=='使用说明':
   ws.cell(1,1,'CASHFLOW｜原生评审工作簿 · v0.6.4');ws.merge_cells('A1:C1');ws['A1'].fill=PatternFill('solid',fgColor=NAVY);ws['A1'].font=Font(name='Microsoft YaHei',size=17,bold=True,color='FFFFFF');ws.row_dimensions[1].height=36;continue
  n=max(1,ws.max_column)
  for row in [1,2,3]:
   if n>1:ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=n)
   c=ws.cell(row,1);c.font=Font(name='Microsoft YaHei',size=17 if row==1 else 10,bold=row==1,color='FFFFFF' if row==1 else '506577');c.fill=PatternFill('solid',fgColor=NAVY if row==1 else 'EAF3F5');c.alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[row].height=34 if row==1 else 30
  for c in ws[4]:c.font=Font(name='Microsoft YaHei',bold=True,color='FFFFFF',size=10);c.fill=PatternFill('solid',fgColor=TEAL);c.alignment=Alignment(wrap_text=True,vertical='center')
  for row in ws.iter_rows(min_row=5):
   for c in row:
    c.font=Font(name='Microsoft YaHei',size=10,color='253B4B');c.alignment=Alignment(wrap_text=True,vertical='top')
    if isinstance(c.value,(int,float)) or (c.data_type=='f' and ws.title=='职业财务摘要'):c.number_format='#,##0.##;[Red]-#,##0.##'
  if ws.title=='职业财务摘要':
   for col in 'ABCDEFGH':ws.column_dimensions[col].width=20
   for i in range(5,17):ws.row_dimensions[i].height=30
  for col,cell in enumerate(ws[4],1):
   if cell.value and any(x in str(cell.value) for x in ['状态','结论']):
    choices='待评审,待执行,待核验,有自动化证据,待用户复验,通过,不通过,有条件通过,阻断,已关闭,重新打开,未签署'
    dv=DataValidation(type='list',formula1='"'+choices+'"');ws.add_data_validation(dv);letter=ws.cell(5,col).column_letter;dv.add(f'{letter}5:{letter}1000');dv.showErrorMessage=False
  ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:4'
 w.save(p)
# Add architecture/flow figures to the native presentation, plus a baseline UI evidence slide.
p=OUT/'00-阅读入口与复核/研发阶段与验收评审简报.pptx';prs=Presentation(p)
for title,image,caption in [('用户回合与子状态',OUT/'03-交互与视觉设计/用户回合流程.png','SVG源文件随包提供；流程图不替代规则条件。'),('系统架构与权威边界',OUT/'03-交互与视觉设计/系统架构.png','本机与在线共享规则，在线权威位于服务端。'),('实现证据：无遮挡查看棋盘',ROOT/'audit/screenshots/v0.6.4/full-board.png','v0.6.4实际截图；页面批准和原版全等仍需独立验收。')]:
 sl=prs.slides.add_slide(prs.slide_layouts[6]);fill=sl.background.fill;fill.solid();fill.fore_color.rgb=RGBColor.from_string('F5F8FB');t=sl.shapes.add_textbox(Inches(.5),Inches(.25),Inches(12),Inches(.6));p0=t.text_frame.paragraphs[0];p0.text=title;p0.font.name='Microsoft YaHei';p0.font.size=Pt(27);p0.font.bold=True;p0.font.color.rgb=RGBColor.from_string(NAVY)
 from PIL import Image
 iw,ih=Image.open(image).size;ratio=min(11.9/iw,5.8/ih);w=iw*ratio;h=ih*ratio;sl.shapes.add_picture(str(image),Inches((13.333-w)/2),Inches(.95),width=Inches(w),height=Inches(h));t=sl.shapes.add_textbox(Inches(.6),Inches(7),Inches(12),Inches(.3));q=t.text_frame.paragraphs[0];q.text=caption;q.font.name='Microsoft YaHei';q.font.size=Pt(11)
prs.save(p)
print('Applied workbook styling, original formulas, chart and three visual review slides.')
